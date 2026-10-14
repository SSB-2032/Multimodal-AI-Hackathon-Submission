import streamlit as st
import pandas as pd
import numpy as np
import re
import os

st.set_page_config(
    page_title="MentorAI — Track D Study Companion",
    page_icon="🎓",
    layout="wide",
    initial_sidebar_state="expanded"
)

st.markdown('''
<style>
    .main-header { font-size: 2.2rem; font-weight: 700; color: #1E3A8A; margin-bottom: 0.2rem; }
    .sub-header { font-size: 1.1rem; color: #4B5563; margin-bottom: 1.5rem; }
    .card { background-color: #F8FAFC; border: 1px solid #E2E8F0; border-radius: 10px; padding: 1.2rem; margin-bottom: 1rem; }
    .citation-badge { background-color: #DBEAFE; color: #1E40AF; padding: 0.2rem 0.5rem; border-radius: 5px; font-weight: 600; font-size: 0.85rem; }
    .metric-box { text-align: center; background-color: #EFF6FF; border-radius: 8px; padding: 1rem; border: 1px solid #BFDBFE; }
    .metric-value { font-size: 1.8rem; font-weight: 700; color: #1D4ED8; }
    .metric-label { font-size: 0.85rem; color: #3B82F6; font-weight: 600; }
</style>
''', unsafe_allow_html=True)

if 'documents' not in st.session_state:
    st.session_state.documents = []

if 'chat_history' not in st.session_state:
    st.session_state.chat_history = []

if 'skills' not in st.session_state:
    st.session_state.skills = {
        'Limit Computation': 65.0,
        'Discontinuity Classification': 45.0,
        '5-Step Continuity Sequence': 50.0
    }

if 'quiz_attempts' not in st.session_state:
    st.session_state.quiz_attempts = []

if 'pre_test' not in st.session_state:
    st.session_state.pre_test = 40.0

if 'post_test' not in st.session_state:
    st.session_state.post_test = 80.0

if 'remediation_event' not in st.session_state:
    st.session_state.remediation_event = None


def extract_text_from_file(file_obj, filename):
    ext = os.path.splitext(filename)[1].lower()
    text = ""
    try:
        if ext in ['.txt', '.md', '.srt', '.vtt']:
            text = file_obj.read().decode('utf-8', errors='ignore')
        elif ext == '.pdf':
            try:
                import pypdf
                reader = pypdf.PdfReader(file_obj)
                pages = [page.extract_text() for page in reader.pages if page.extract_text()]
                text = "\n\n".join(pages)
            except Exception as e:
                text = f"[PDF Parsing Error for {filename}: {e}]"
        elif ext == '.docx':
            try:
                import docx
                doc = docx.Document(file_obj)
                paras = [p.text for p in doc.paragraphs if p.text.strip()]
                text = "\n\n".join(paras)
            except Exception as e:
                text = f"[DOCX Parsing Error for {filename}: {e}]"
        elif ext == '.pptx':
            try:
                import pptx
                prs = pptx.Presentation(file_obj)
                slides_text = []
                for i, slide in enumerate(prs.slides):
                    slide_content = [shape.text.strip() for shape in slide.shapes if hasattr(shape, 'text') and shape.text.strip()]
                    if slide_content:
                        slides_text.append(f"--- Slide {i+1} ---\n" + "\n".join(slide_content))
                text = "\n\n".join(slides_text)
            except Exception as e:
                text = f"[PPTX Parsing Error for {filename}: {e}]"
        elif ext == '.csv':
            try:
                df = pd.read_csv(file_obj)
                text = df.to_string()
            except Exception as e:
                text = f"[CSV Parsing Error for {filename}: {e}]"
        elif ext == '.json':
            try:
                import json
                data = json.load(file_obj)
                text = json.dumps(data, indent=2)
            except Exception as e:
                text = f"[JSON Parsing Error for {filename}: {e}]"
        else:
            text = file_obj.read().decode('utf-8', errors='ignore')
    except Exception as e:
        text = f"[File Reading Error: {e}]"
    return text

def load_bundled_demo():
    demo_file = os.path.join(os.path.dirname(__file__), 'demo_course.txt')
    if os.path.exists(demo_file):
        with open(demo_file, 'r', encoding='utf-8') as f:
            content = f.read()
        chunks = [c.strip() for c in content.split('\n\n') if c.strip()]
        st.session_state.documents = [
            {'id': i + 1, 'source': 'Calculus_Demo_Course.txt', 'chunk': chunk}
            for i, chunk in enumerate(chunks)
        ]
        return len(chunks)
    return 0

def retrieve_evidence(query, top_k=3):
    if not st.session_state.documents:
        return []
    query_words = set(re.findall(r'\w+', query.lower()))
    scored_docs = []
    for doc in st.session_state.documents:
        doc_words = set(re.findall(r'\w+', doc['chunk'].lower()))
        overlap = len(query_words.intersection(doc_words))
        score = overlap / (len(query_words) + 1e-5)
        scored_docs.append((score, doc))
    scored_docs.sort(key=lambda x: x[0], reverse=True)
    return [doc for score, doc in scored_docs[:top_k] if score > 0] or [st.session_state.documents[0]]

def generate_answer(query, retrieved_chunks, api_key=None):
    if api_key:
        try:
            from openai import OpenAI
            client = OpenAI(api_key=api_key)
            context_text = '\n\n'.join([f"[Source {d['id']} ({d['source']})]: {d['chunk']}" for d in retrieved_chunks])
            prompt = f"You are MentorAI, a source-cited tutor. Answer the user question using ONLY the provided course context. Cite sources as [Source X].\n\nContext:\n{context_text}\n\nQuestion: {query}"
            response = client.chat.completions.create(
                model='gpt-4o-mini',
                messages=[{'role': 'user', 'content': prompt}],
                max_tokens=400
            )
            return response.choices[0].message.content
        except Exception as e:
            st.warning(f"API Call failed, falling back to local deterministic generator: {e}")

    context_summary = ' '.join([d['chunk'] for d in retrieved_chunks])
    citations = ', '.join([f"[Source {d['id']}]" for d in retrieved_chunks])
    return f"Based on your course materials {citations}:\n\n{context_summary}\n\n*(Grounded directly in course evidence)*"

with st.sidebar:
    st.title("MentorAI Control Panel")
    st.caption("Personalised Tutoring & Adaptive Learning")
    
    st.subheader("1. Quick Preset")
    if st.button("⚡ Load Calculus Demo Course", use_container_width=True, type="primary"):
        n_chunks = load_bundled_demo()
        st.success(f"Loaded {n_chunks} evidence chunks!")
        
    st.divider()
    
    st.subheader("2. Upload Course Materials")
    uploaded_files = st.file_uploader(
        "Ingest Textbooks, Slides, Transcripts",
        type=["pdf", "docx", "pptx", "csv", "json", "txt", "md", "srt", "vtt"],
        accept_multiple_files=True
    )
    if uploaded_files:
        new_docs = []
        for f in uploaded_files:
            text = extract_text_from_file(f, f.name)
            ext_label = os.path.splitext(f.name)[1].upper().replace('.', '')
            chunks = [c.strip() for c in text.split('\n\n') if c.strip()]
            for chunk in chunks:
                new_docs.append({'id': len(st.session_state.documents) + len(new_docs) + 1, 'source': f.name, 'chunk': chunk, 'format': ext_label})
        st.session_state.documents.extend(new_docs)
        st.success(f"Added {len(new_docs)} new chunks!")

    st.divider()
    
    st.subheader("3. API Key (Optional)")
    api_key_input = st.text_input("OpenAI API Key", help="Optional. System falls back to deterministic local generation if omitted.")
    
    st.divider()

st.markdown("<div class='main-header'>MentorAI: Personalised AI Study Companion</div>", unsafe_allow_html=True)
st.markdown("<div class='sub-header'>Track D Problem Statement | Source-Cited Knowledge Base & Adaptive Tutoring</div>", unsafe_allow_html=True)

tab_tutor, tab_assess, tab_gain = st.tabs([
    "📚 Source-Cited Tutor", 
    "🎯 Adaptive Assessment & Diagnosis", 
    "📊 Measurable Learning Gain"
])

with tab_tutor:
    st.subheader("Course Knowledge Base & Grounded Q&A")
    
    if not st.session_state.documents:
        st.warning("⚠️ No course materials loaded yet. Click '⚡ Load Calculus Demo Course' in the sidebar to begin!")
    else:
        st.success(f"Active Knowledge Base: **{len(st.session_state.documents)} chunks** indexed across course files.")
        
    query_input = st.text_input("Ask MentorAI a question from your course:", value="What is a removable discontinuity and how do I solve it?")
    
    if st.button("Ask Tutor", type="primary"):
        if not st.session_state.documents:
            load_bundled_demo()
            
        with st.spinner("Retrieving course evidence and synthesizing response..."):
            retrieved = retrieve_evidence(query_input, top_k=3)
            answer = generate_answer(query_input, retrieved, api_key_input)
            
            st.session_state.chat_history.append({
                "query": query_input,
                "answer": answer,
                "evidence": retrieved
            })
            
    for chat in reversed(st.session_state.chat_history):
        with st.container():
            st.markdown(f"**🙋 Student Query**: {chat['query']}")
            st.markdown(f"**🤖 MentorAI Answer**:\n{chat['answer']}")
            
            with st.expander("🔍 View Retrieved Course Evidence Chunks"):
                for doc in chat['evidence']:
                    st.markdown(f"<span class='citation-badge'>Source {doc['id']} ({doc['source']})</span>", unsafe_allow_html=True)
                    st.caption(doc['chunk'])
            st.divider()

with tab_assess:
    st.subheader("Adaptive Assessment & Skill Diagnosis Engine")
    
    col_left, col_right = st.columns([3, 2])
    
    with col_right:
        st.markdown("<div class='card'>", unsafe_allow_html=True)
        st.markdown("### 📊 Explainable Learner State")
        st.caption("Granular Skill Mastery Model")
        
        for skill_name, score in st.session_state.skills.items():
            st.write(f"**{skill_name}**")
            st.progress(score / 100.0)
            st.caption(f"Mastery: **{score:.1f}%**")
            
        st.markdown("</div>", unsafe_allow_html=True)
        
        if st.session_state.remediation_event:
            st.error("🚨 **Targeted Remediation Triggered!**")
            st.markdown(f"**Diagnosed Weak Skill**: `{st.session_state.remediation_event['skill']}`")
            st.warning(f"**Remedial Micro-Lesson**: {st.session_state.remediation_event['explanation']}")
            
    with col_left:
        st.markdown("### 📝 Dynamic Assessment")
        
        questions = [
            {
                "id": 1,
                "question": r"For $f(x) = \frac{x^2 - 1}{x - 1}$ at $x = 1$, what type of discontinuity is present?",
                "options": ["Jump Discontinuity", "Removable Discontinuity", "Infinite Discontinuity", "No Discontinuity"],
                "correct": "Removable Discontinuity",
                "skill": "Discontinuity Classification",
                "explanation": r"Because $\lim_{x \to 1} f(x) = 2$ exists but $f(1)$ is undefined, the discontinuity is removable."
            },
            {
                "id": 2,
                "question": "What is the 3rd step in the recommended 5-step continuity evaluation sequence?",
                "options": ["Check if $f(a)$ is defined", r"Compute left-hand limit $\lim_{x \to a^-} f(x)$", r"Compute right-hand limit $\lim_{x \to a^+} f(x)$", "Compare limits"],
                "correct": r"Compute right-hand limit $\lim_{x \to a^+} f(x)$",
                "skill": "5-Step Continuity Sequence",
                "explanation": r"Step 1: check $f(a)$, Step 2: left-hand limit $\lim_{x \to a^-} f(x)$, Step 3: right-hand limit $\lim_{x \to a^+} f(x)$."
            },
            {
                "id": 3,
                "question": r"If $\lim_{x \to a^-} f(x) = 3$ and $\lim_{x \to a^+} f(x) = 5$, what can be concluded?",
                "options": ["Removable Discontinuity", "Jump Discontinuity", "Continuous at $x=a$", "Undefined Limit"],
                "correct": "Jump Discontinuity",
                "skill": "Limit Computation",
                "explanation": r"When left-hand and right-hand limits exist but are unequal ($\lim_{x \to a^-} f(x) \neq \lim_{x \to a^+} f(x)$), a jump discontinuity exists."
            }
        ]
        
        q_idx = len(st.session_state.quiz_attempts) % len(questions)
        q = questions[q_idx]
        
        with st.container(border=True):
            st.markdown(f"**Question {q['id']}**: {q['question']}")
            st.caption(f"Targeted Skill: `{q['skill']}`")
            
            selected_option = st.radio("Select your answer:", q["options"], key=f"q_{q['id']}")
            
            if st.button("Submit Answer", type="primary"):
                is_correct = (selected_option == q["correct"])
                
                current_skill_val = st.session_state.skills[q["skill"]]
                if is_correct:
                    st.session_state.skills[q["skill"]] = min(100.0, current_skill_val + 15.0)
                    st.session_state.remediation_event = None
                    st.success("🎉 **Correct!** Your skill mastery for " + q["skill"] + " increased.")
                else:
                    st.session_state.skills[q["skill"]] = max(0.0, current_skill_val - 20.0)
                    st.session_state.remediation_event = {
                        "skill": q["skill"],
                        "explanation": q["explanation"]
                    }
                    st.error("❌ **Incorrect.** " + q["explanation"])
                    
                st.session_state.quiz_attempts.append({
                    "question": q["question"],
                    "selected": selected_option,
                    "correct": is_correct,
                    "skill": q["skill"]
                })

with tab_gain:
    st.subheader("Measurable Learning Outcome Engine (g)")
    st.markdown("MentorAI tracks student growth using **Hakes Normalized Learning Gain Formula**:")
    st.latex(r"g = \frac{\text{Post\%} - \text{Pre\%}}{100\% - \text{Pre\%}}")
    
    col1, col2, col3 = st.columns(3)
    with col1:
        st.session_state.pre_test = st.number_input("Diagnostic Pre-Test Score (%)", min_value=0.0, max_value=100.0, value=float(st.session_state.pre_test))
    with col2:
        st.session_state.post_test = st.number_input("Post-Intervention Score (%)", min_value=0.0, max_value=100.0, value=float(st.session_state.post_test))
        
    pre = st.session_state.pre_test
    post = st.session_state.post_test
    abs_gain = post - pre
    norm_gain = (post - pre) / (100.0 - pre) if pre < 100 else 0.0
    
    with col3:
        st.markdown("<div class='metric-box'>", unsafe_allow_html=True)
        st.markdown(f"<div class='metric-value'>{norm_gain:.2f}</div>", unsafe_allow_html=True)
        st.markdown("<div class='metric-label'>Normalized Gain (g)</div>", unsafe_allow_html=True)
        st.markdown("</div>", unsafe_allow_html=True)
        
    st.divider()
    col_chart, col_summary = st.columns([2, 1])
    with col_chart:
        df_chart = pd.DataFrame({
            "Evaluation Phase": ["Diagnostic Pre-Test", "Post-Intervention Test"],
            "Score (%)": [pre, post]
        })
        st.bar_chart(df_chart, x="Evaluation Phase", y="Score (%)", color="#1D4ED8")
        
    with col_summary:
        st.markdown("<div class='card'>", unsafe_allow_html=True)
        st.markdown("### 📈 Impact Summary")
        st.write(f"* **Absolute Gain**: +{abs_gain:.1f}%")
        st.write(f"* **Normalized Gain (g)**: {norm_gain:.2f}")
        if norm_gain >= 0.7:
            st.success("🌟 **Category**: High Gain (g >= 0.7)")
        elif norm_gain >= 0.3:
            st.info("👍 **Category**: Medium Gain (0.3 <= g < 0.7)")
        else:
            st.warning("⚠️ **Category**: Low Gain (g < 0.3)")
        st.markdown("</div>", unsafe_allow_html=True)

st.caption("MentorAI — Personalised AI Study Companion")
